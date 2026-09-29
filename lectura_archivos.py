"""Lectura segura de enunciados para la Calculadora de Tuberías de Mott.

V14.2
- TXT: UTF-8 / Windows-1252 / Latin-1.
- DOCX: párrafos y tablas.
- PDF con texto seleccionable: extracción con pypdf.
- PNG / JPG / JPEG: OCR con Tesseract.
- PDF escaneado: detección automática y OCR por página con PyMuPDF + Tesseract.

La salida SIEMPRE se entrega como texto editable para revisión antes de enviarla
al clasificador. V14.2 no interpreta todavía diagramas hidráulicos como estructura
geométrica ni resuelve fórmulas manuscritas a partir de la imagen.
"""

from __future__ import annotations

from dataclasses import dataclass, asdict
from io import BytesIO
from pathlib import Path
import os
import re
import shutil

MAX_ARCHIVO_BYTES = 20 * 1024 * 1024  # 20 MB
EXTENSIONES_SOPORTADAS = (".txt", ".docx", ".pdf", ".png", ".jpg", ".jpeg")


class ErrorLecturaArchivo(ValueError):
    """Error controlado al leer un archivo de enunciado."""


@dataclass
class ResultadoLectura:
    texto: str
    nombre: str
    extension: str
    caracteres: int
    paginas: int | None = None
    advertencias: list[str] | None = None
    metodo: str = ""
    ocr_aplicado: bool = False

    def como_dict(self):
        d = asdict(self)
        d["advertencias"] = list(self.advertencias or [])
        return d


def _limpiar_texto(texto: str) -> str:
    texto = (texto or "").replace("\x00", "")
    texto = texto.replace("\r\n", "\n").replace("\r", "\n")
    texto = re.sub(r"\n[ \t]+\n", "\n\n", texto)
    texto = re.sub(r"\n{4,}", "\n\n\n", texto)
    return texto.strip()


def _validar(nombre: str, datos: bytes) -> str:
    if not nombre:
        raise ErrorLecturaArchivo("El archivo no tiene nombre identificable.")

    ext = Path(nombre).suffix.lower()
    if ext not in EXTENSIONES_SOPORTADAS:
        raise ErrorLecturaArchivo(
            f"Formato no admitido en V14.2: {ext or 'sin extensión'}. "
            "Use TXT, DOCX, PDF, PNG, JPG o JPEG."
        )

    if not datos:
        raise ErrorLecturaArchivo("El archivo está vacío.")

    if len(datos) > MAX_ARCHIVO_BYTES:
        raise ErrorLecturaArchivo(
            f"El archivo supera el límite de {MAX_ARCHIVO_BYTES // (1024 * 1024)} MB de V14.2."
        )

    return ext


# ============================================================
# OCR
# ============================================================


def _candidatos_tesseract():
    candidatos = []

    env_cmd = os.environ.get("TESSERACT_CMD", "").strip()
    if env_cmd:
        candidatos.append(env_cmd)

    which_cmd = shutil.which("tesseract")
    if which_cmd:
        candidatos.append(which_cmd)

    candidatos.extend(
        [
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        ]
    )
    return candidatos


def ocr_disponible() -> tuple[bool, str]:
    """Comprueba bibliotecas Python y el ejecutable Tesseract."""
    try:
        import pytesseract  # noqa: F401
    except ImportError:
        return False, "Falta pytesseract. Instale con: python -m pip install pytesseract"

    try:
        from PIL import Image  # noqa: F401
    except ImportError:
        return False, "Falta Pillow. Instale con: python -m pip install pillow"

    for ruta in _candidatos_tesseract():
        if ruta and Path(ruta).exists():
            return True, str(ruta)

    return False, (
        "No se encontró Tesseract OCR. Instale Tesseract para Windows. "
        r"La ruta habitual es C:\Program Files\Tesseract-OCR\tesseract.exe. "
        "Si lo instaló en otra ubicación, defina TESSERACT_CMD con esa ruta."
    )


def _asegurar_ocr():
    disponible, detalle = ocr_disponible()
    if not disponible:
        raise ErrorLecturaArchivo(detalle)

    import pytesseract

    pytesseract.pytesseract.tesseract_cmd = detalle
    return pytesseract


def _preprocesar_imagen_pil(img):
    from PIL import ImageOps

    # EXIF puede contener la orientación correcta de fotos hechas con celular.
    try:
        img = ImageOps.exif_transpose(img)
    except Exception:
        pass

    img = img.convert("L")
    img = ImageOps.autocontrast(img)

    # Si la imagen es pequeña, escalar ayuda al OCR; si ya es grande, evita inflarla demasiado.
    lado_mayor = max(img.size)
    if lado_mayor < 2200:
        escala = min(2.5, 2200.0 / max(lado_mayor, 1))
        nuevo = (max(1, int(img.width * escala)), max(1, int(img.height * escala)))
        img = img.resize(nuevo)

    return img


def _ocr_desde_pil(img) -> str:
    pytesseract = _asegurar_ocr()
    img = _preprocesar_imagen_pil(img)

    # Prioriza español+inglés. Si no está instalado el idioma spa, cae a eng/default.
    intentos = [
        ("spa+eng", "--psm 6"),
        ("eng", "--psm 6"),
        (None, "--psm 6"),
    ]

    errores = []
    for idioma, config in intentos:
        try:
            kwargs = {"config": config}
            if idioma:
                kwargs["lang"] = idioma
            texto = pytesseract.image_to_string(img, **kwargs)
            texto = _limpiar_texto(texto)
            if texto:
                return texto
        except Exception as e:
            errores.append(str(e))

    if errores:
        raise ErrorLecturaArchivo(f"No se pudo completar el OCR: {errores[-1]}")
    return ""


# ============================================================
# Lectores por formato
# ============================================================


def _extraer_txt(datos: bytes):
    advertencias = []
    intentos = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
    ultimo_error = None

    for codificacion in intentos:
        try:
            texto = datos.decode(codificacion)
            if codificacion not in ("utf-8-sig", "utf-8"):
                advertencias.append(
                    f"El TXT se decodificó como {codificacion}; revise tildes, símbolos y unidades."
                )
            return texto, advertencias, f"TXT / {codificacion}", False
        except UnicodeDecodeError as e:
            ultimo_error = e

    raise ErrorLecturaArchivo(f"No se pudo decodificar el TXT: {ultimo_error}")


def _extraer_docx(datos: bytes):
    try:
        from docx import Document
    except ImportError as e:
        raise ErrorLecturaArchivo(
            "Falta python-docx. Instale con: python -m pip install python-docx"
        ) from e

    try:
        doc = Document(BytesIO(datos))
    except Exception as e:
        raise ErrorLecturaArchivo(f"No se pudo abrir el DOCX: {e}") from e

    bloques = []

    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            bloques.append(t)

    for tabla in doc.tables:
        for fila in tabla.rows:
            celdas = [celda.text.strip() for celda in fila.cells]
            if any(celdas):
                bloques.append("\t".join(celdas))

    advertencias = [
        "V14.2 extrae texto y tablas del DOCX, pero todavía no interpreta imágenes, "
        "diagramas ni contenido escaneado embebido dentro del documento."
    ]
    return "\n".join(bloques), advertencias, "DOCX / python-docx", False


def _extraer_imagen(datos: bytes, ext: str):
    try:
        from PIL import Image
    except ImportError as e:
        raise ErrorLecturaArchivo(
            "Falta Pillow. Instale con: python -m pip install pillow"
        ) from e

    try:
        img = Image.open(BytesIO(datos))
        img.load()
    except Exception as e:
        raise ErrorLecturaArchivo(f"No se pudo abrir la imagen: {e}") from e

    texto = _ocr_desde_pil(img)
    advertencias = [
        "V14.2 aplicó OCR a la imagen. Revise el texto antes de analizar: números, unidades, "
        "grados, símbolos de diámetro y superíndices pueden requerir corrección manual.",
        "Todavía no se interpretan diagramas hidráulicos ni fórmulas manuscritas como estructura matemática.",
    ]
    return texto, 1, advertencias, f"Imagen {ext.upper().lstrip('.')} / OCR Tesseract", True


def _ocr_pdf_escaneado(datos: bytes):
    try:
        import fitz
    except ImportError as e:
        raise ErrorLecturaArchivo(
            "Falta PyMuPDF. Instale con: python -m pip install pymupdf"
        ) from e

    try:
        from PIL import Image
    except ImportError as e:
        raise ErrorLecturaArchivo(
            "Falta Pillow. Instale con: python -m pip install pillow"
        ) from e

    try:
        doc = fitz.open(stream=datos, filetype="pdf")
    except Exception as e:
        raise ErrorLecturaArchivo(f"No se pudo abrir el PDF escaneado: {e}") from e

    bloques = []
    try:
        for i, pagina in enumerate(doc, 1):
            pix = pagina.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            modo = "RGB" if pix.n >= 3 else "L"
            img = Image.frombytes(modo, [pix.width, pix.height], pix.samples)
            texto_pag = _ocr_desde_pil(img)
            if texto_pag.strip():
                bloques.append(f"[Página {i}]\n{texto_pag.strip()}")
    finally:
        doc.close()

    return "\n\n".join(bloques), len(bloques)


def _extraer_pdf(datos: bytes):
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ErrorLecturaArchivo(
            "Falta pypdf. Instale con: python -m pip install pypdf"
        ) from e

    try:
        reader = PdfReader(BytesIO(datos))
    except Exception as e:
        raise ErrorLecturaArchivo(f"No se pudo abrir el PDF: {e}") from e

    paginas = []
    paginas_sin_texto = 0

    for pagina in reader.pages:
        try:
            t = pagina.extract_text() or ""
        except Exception:
            t = ""
        t = t.strip()
        if not t:
            paginas_sin_texto += 1
        paginas.append(t)

    texto_pdf = _limpiar_texto("\n\n".join(t for t in paginas if t))
    total_paginas = len(reader.pages)

    # PDF totalmente escaneado o casi sin texto útil.
    usar_ocr = (
        not texto_pdf
        or (
            len(texto_pdf) < 80
            and paginas_sin_texto >= max(1, (total_paginas + 1) // 2)
        )
    )

    if usar_ocr:
        texto_ocr, paginas_con_ocr = _ocr_pdf_escaneado(datos)
        advertencias = [
            "No se detectó texto seleccionable suficiente; se aplicó OCR automáticamente al PDF escaneado.",
            "Revise el texto extraído por OCR antes de analizar: símbolos, superíndices, diámetros y ecuaciones pueden requerir corrección manual.",
        ]
        return (
            texto_ocr,
            total_paginas,
            advertencias,
            "PDF escaneado / OCR Tesseract + PyMuPDF",
            True,
        )

    advertencias = [
        "Revise el texto extraído del PDF antes de analizarlo: columnas, ecuaciones y símbolos pueden cambiar de orden."
    ]
    if paginas_sin_texto:
        advertencias.append(
            f"{paginas_sin_texto} de {total_paginas} página(s) no devolvieron texto seleccionable."
        )

    return texto_pdf, total_paginas, advertencias, "PDF / pypdf", False


def extraer_texto_archivo(nombre: str, datos: bytes) -> ResultadoLectura:
    """Extrae texto y devuelve siempre una versión editable antes del clasificador."""
    ext = _validar(nombre, datos)

    if ext == ".txt":
        texto, advertencias, metodo, ocr_aplicado = _extraer_txt(datos)
        paginas = None
    elif ext == ".docx":
        texto, advertencias, metodo, ocr_aplicado = _extraer_docx(datos)
        paginas = None
    elif ext == ".pdf":
        texto, paginas, advertencias, metodo, ocr_aplicado = _extraer_pdf(datos)
    elif ext in (".png", ".jpg", ".jpeg"):
        texto, paginas, advertencias, metodo, ocr_aplicado = _extraer_imagen(datos, ext)
    else:  # pragma: no cover
        raise ErrorLecturaArchivo(f"Formato no manejado internamente: {ext}")

    texto = _limpiar_texto(texto)
    if not texto:
        raise ErrorLecturaArchivo(
            "No se obtuvo texto utilizable del archivo. Si la imagen o PDF escaneado está "
            "borroso, inclinado o tiene poca resolución, el OCR puede no reconocerlo correctamente."
        )

    return ResultadoLectura(
        texto=texto,
        nombre=nombre,
        extension=ext,
        caracteres=len(texto),
        paginas=paginas,
        advertencias=advertencias,
        metodo=metodo,
        ocr_aplicado=ocr_aplicado,
    )
