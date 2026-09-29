"""Reporte técnico V2 — V14.18.

Capa de presentación profesional construida encima del reporte trazable V14.9.
No modifica resultados, auditorías ni ecuaciones del solver.
"""
from __future__ import annotations

import io
import math
from datetime import datetime

VERSION_REPORTE_V2 = "V14.18"


class ErrorReporteV1418(ValueError):
    pass


def _fmt(v, nd=6):
    if v is None:
        return "—"
    if isinstance(v, bool):
        return "Sí" if v else "No"
    if isinstance(v, float):
        if not math.isfinite(v):
            return str(v)
        av = abs(v)
        if av != 0 and (av < 1e-4 or av >= 1e6):
            return f"{v:.5e}"
        return f"{v:.{nd}g}"
    return str(v)


def _resultado_principal(reporte: dict) -> tuple[str, str]:
    clase = str(reporte.get("clase") or "")
    r = reporte.get("resultado") or {}
    if clase.startswith("Clase II"):
        return "Caudal Q", f"{_fmt(r.get('Q_m3s'), 8)} m³/s"
    if clase == "Clase III-A":
        d = r.get("D_mott") or r.get("D_m") or r.get("diametro_m")
        return "Diámetro mínimo", f"{_fmt(d, 8)} m"
    if clase == "Clase III-B":
        p = r.get("P2_calculada")
        if p is not None:
            return "Presión P2 calculada", f"{_fmt(float(p)/1000.0, 8)} kPa"
    if clase == "Clase I":
        return str(r.get("incognita") or "Resultado"), f"{_fmt(r.get('valor'), 8)} {r.get('unidad') or ''}".strip()
    return "Resultado", "Ver sección de resultados"


def _ecuaciones_metodo(clase: str, resultado: dict | None = None) -> list[tuple[str, str]]:
    comunes = [
        ("Continuidad", "A = πD²/4 ; V = Q/A"),
        ("Reynolds", "Re = VD/ν"),
        ("Darcy-Weisbach", "hf = f(L/D)(V²/2g)"),
        ("Pérdidas menores", "hm = ΣK(V²/2g)"),
        ("Ecuación general de energía", "P1/γ + z1 + V1²/(2g) + hA - hR - hL = P2/γ + z2 + V2²/(2g)"),
    ]
    if clase in ("Clase II-B", "Clase II-C"):
        comunes.append(("Cierre de caudal", "Se busca Q tal que el residual de la ecuación de energía sea aproximadamente cero."))
    elif clase == "Clase II-A":
        metodo = str((resultado or {}).get("metodo_calculo_principal") or "")
        if "prueba y error" in metodo.lower():
            comunes.append(("Método II-A solicitado", "Prueba y error con Darcy-Weisbach: se itera el caudal hasta cerrar la ecuación de energía; la Ec. (11-3) puede usarse solo como estimación inicial."))
        else:
            comunes.append(("Método II-A de Mott", "Se aplica la forma explícita correspondiente de Mott y se contrasta con una verificación numérica independiente."))
    elif clase == "Clase III-A":
        comunes.append(("Método III-A de Mott", "Se obtiene el diámetro mínimo mediante la ecuación explícita de Mott y se verifica con Colebrook."))
    elif clase == "Clase III-B":
        comunes.append(("Verificación III-B", "Con un diámetro comercial dado se calcula la presión disponible y se compara con la mínima requerida."))
    return comunes


def _supuestos(reporte: dict) -> list[str]:
    entradas = reporte.get("entradas") or {}
    generales = {x.get("clave"): x for x in entradas.get("generales") or []}
    sup = [
        "Régimen permanente y análisis unidimensional con velocidad media por tramo.",
        "Las pérdidas distribuidas se evalúan con Darcy-Weisbach y las pérdidas locales con los modelos/K confirmados en el expediente.",
        "Las propiedades, materiales, accesorios y transiciones usados son los confirmados por el flujo trazable; el reporte no sustituye datos faltantes por valores inventados.",
    ]
    for clave, etiqueta in (("P2_kpa", "P2"), ("hA_m", "hA"), ("hR_m", "hR")):
        x = generales.get(clave)
        if x and "inferencia segura" in str(x.get("fuente") or "").lower():
            sup.append(f"{etiqueta} fue incorporado mediante una inferencia segura V14.5 y quedó registrado como tal en la trazabilidad.")
    if reporte.get("geometria_v1414"):
        sup.append("La evidencia V14.14 se usa solo como geometría relativa; no convierte píxeles a longitudes o elevaciones físicas.")
    return sup


def _nomenclatura() -> list[tuple[str, str, str]]:
    return [
        ("Q", "Caudal volumétrico", "m³/s"),
        ("D", "Diámetro interior", "m"),
        ("V", "Velocidad media", "m/s"),
        ("Re", "Número de Reynolds", "—"),
        ("f", "Factor de fricción de Darcy", "—"),
        ("ε", "Rugosidad absoluta", "m"),
        ("hf", "Pérdida distribuida", "m de fluido"),
        ("hm", "Pérdida menor/local", "m de fluido"),
        ("hL", "Pérdida total", "m de fluido"),
        ("LE / EGL", "Línea de energía", "m de fluido"),
        ("LAM / HGL", "Línea piezométrica", "m de fluido"),
        ("γ", "Peso específico", "N/m³"),
        ("ν", "Viscosidad cinemática", "m²/s"),
    ]


def _figuras_png(contexto: dict | None) -> list[tuple[str, bytes]]:
    if not contexto or not contexto.get("disponible"):
        return []
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from diagramas import crear_esquema_sistema, crear_lineas_energia, crear_perdidas_acumuladas_v1415
    except Exception as e:
        raise ErrorReporteV1418(f"No se pudieron cargar los módulos gráficos V14.15: {e}") from e

    sis = contexto.get("sistema") or {}
    tramos = contexto.get("tramos") or []
    resultados = sis.get("resultados") or []
    trans = contexto.get("transiciones") or []
    figs = []
    specs = []
    try:
        specs.append(("Esquema hidráulico", crear_esquema_sistema(
            tramos=tramos,
            z1=contexto["z1_m"], z2=contexto["z2_m"],
            tipo_v1=contexto.get("tipo_v1", "tuberia"), tipo_v2=contexto.get("tipo_v2", "tuberia"),
            hA=contexto.get("hA_m", 0.0), hR=contexto.get("hR_m", 0.0),
            posicion_bomba=contexto.get("posicion_bomba_m"), posicion_turbina=contexto.get("posicion_turbina_m"),
            z_nodos=contexto.get("z_nodos_m") or None, transiciones=trans,
            titulo="Esquema hidráulico del sistema",
            mostrar_etiquetas_tramos=True, mostrar_etiquetas_accesorios=True,
            mostrar_etiquetas_transiciones=True, mostrar_etiquetas_equipos=True, mostrar_sentido_flujo=True,
        )))
        specs.append(("Línea de energía y línea piezométrica", crear_lineas_energia(
            resultados=resultados, gamma=contexto["gamma"],
            P1=contexto["P1_pa"], z1=contexto["z1_m"], P2=contexto["P2_pa"], z2=contexto["z2_m"],
            V1=contexto["V1_ms"], V2=contexto["V2_ms"],
            hA=contexto.get("hA_m", 0.0), hR=contexto.get("hR_m", 0.0), tramos=tramos,
            posicion_bomba=contexto.get("posicion_bomba_m"), posicion_turbina=contexto.get("posicion_turbina_m"),
            transiciones=trans, titulo="LE / LAM del sistema",
            mostrar_marcadores_eventos=True, mostrar_etiquetas_eventos=False,
        )))
        specs.append(("Distribución acumulada de pérdidas", crear_perdidas_acumuladas_v1415(
            resultados=resultados, tramos=tramos, transiciones=trans,
            hA=contexto.get("hA_m", 0.0), hR=contexto.get("hR_m", 0.0),
            posicion_bomba=contexto.get("posicion_bomba_m"), posicion_turbina=contexto.get("posicion_turbina_m"),
            mostrar_etiquetas_eventos=False, titulo="Pérdidas acumuladas del sistema",
        )))
        for titulo, fig in specs:
            bio = io.BytesIO()
            fig.savefig(bio, format="png", dpi=180, bbox_inches="tight")
            plt.close(fig)
            figs.append((titulo, bio.getvalue()))
    finally:
        for _, fig in specs:
            try:
                plt.close(fig)
            except Exception:
                pass
    return figs


def reporte_docx_v2(reporte: dict, contexto_diagramas: dict | None = None) -> bytes:
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.section import WD_SECTION, WD_ORIENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as e:
        raise ErrorReporteV1418("Falta python-docx para generar el reporte Word V14.18.") from e

    r = reporte or {}
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.7); sec.right_margin = Inches(0.7)
    portrait_width, portrait_height = sec.page_width, sec.page_height

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    for s in ("Title", "Heading 1", "Heading 2"):
        styles[s].font.name = "Aptos Display"

    # Portada
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.space_before = Pt(65)
    rr = p.add_run("REPORTE TÉCNICO DE SISTEMA DE TUBERÍAS")
    rr.bold = True; rr.font.size = Pt(22)
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p2.add_run(str(r.get("clase") or "Sistema de tuberías")); rr.bold = True; rr.font.size = Pt(16)
    p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p3.add_run("Método de Mott · cálculo, auditoría y trazabilidad"); rr.font.size = Pt(12)
    for _ in range(6): doc.add_paragraph("")
    titulo_res, valor_res = _resultado_principal(r)
    t = doc.add_table(rows=4, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
    portada = [
        ("Resultado principal", f"{titulo_res}: {valor_res}"),
        ("Auditoría matemática", str(r.get("estado_matematico") or "—")),
        ("Auditoría física", str(r.get("estado_fisico") or "—")),
        ("Trazabilidad", f"Snapshot V14.7: {str(r.get('firma_transferencia_v147') or '—')[:12]}"),
    ]
    for row, (a,b) in zip(t.rows, portada):
        row.cells[0].text = a; row.cells[1].text = b
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Generado por {VERSION_REPORTE_V2} · {datetime.now().strftime('%Y-%m-%d %H:%M')}").italic = True
    doc.add_page_break()

    # Resumen
    doc.add_heading("1. Resumen técnico", level=1)
    doc.add_paragraph(
        f"Se resolvió un problema de {r.get('clase','—')} mediante el motor hidráulico validado de la calculadora. "
        f"El resultado principal es {titulo_res.lower()} = {valor_res}. "
        f"La auditoría matemática quedó en {r.get('estado_matematico','—')} y la auditoría física en {r.get('estado_fisico','—')}."
    )
    doc.add_heading("2. Enunciado", level=1)
    doc.add_paragraph(r.get("enunciado") or "No disponible en el manifiesto.")

    # Datos
    doc.add_heading("3. Datos, condiciones y procedencia", level=1)
    gens = (r.get("entradas") or {}).get("generales") or []
    table = doc.add_table(rows=1, cols=4); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(("Dato","Valor","Unidad","Fuente")): table.rows[0].cells[i].text = h
    for x in gens:
        c = table.add_row().cells
        vals = (x.get("dato",""), _fmt(x.get("valor")), x.get("unidad",""), x.get("fuente",""))
        for i,v in enumerate(vals): c[i].text = str(v)

    doc.add_heading("4. Tramos, accesorios y transiciones", level=1)
    tramos = (r.get("entradas") or {}).get("tramos") or []
    tab = doc.add_table(rows=1, cols=6); tab.style = "Table Grid"; tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(("Tramo","L (m)","D (m)","Material","ε (m)","Accesorios")): tab.rows[0].cells[i].text = h
    for t0 in tramos:
        acc = ", ".join(f"{a.get('cantidad',1)}× {a.get('nombre','Accesorio')}" for a in t0.get("accesorios") or []) or "—"
        vals = (t0.get("numero"), _fmt(t0.get("L_m")), _fmt(t0.get("D_m")), t0.get("material","—"), _fmt(t0.get("epsilon_m")), acc)
        c = tab.add_row().cells
        for i,v in enumerate(vals): c[i].text = str(v)
    trans = (r.get("entradas") or {}).get("transiciones") or []
    if trans:
        doc.add_paragraph("Transiciones confirmadas:")
        for x in trans:
            doc.add_paragraph(f"Entre tramos {x.get('entre','—')} y {int(x.get('entre',0) or 0)+1}: {x.get('tipo','—')}", style="List Bullet")

    doc.add_heading("5. Supuestos y convenciones", level=1)
    for s in _supuestos(r): doc.add_paragraph(s, style="List Bullet")

    doc.add_heading("6. Ecuaciones y método de solución", level=1)
    for nombre, ecu in _ecuaciones_metodo(str(r.get("clase") or ""), r.get("resultado") or {}):
        p = doc.add_paragraph()
        p.add_run(nombre + ": ").bold = True
        p.add_run(ecu)

    # Resultado principal y balance
    doc.add_heading("7. Resultados", level=1)
    result = r.get("resultado") or {}
    rt = doc.add_table(rows=1, cols=2); rt.style = "Table Grid"; rt.alignment = WD_TABLE_ALIGNMENT.CENTER
    rt.rows[0].cells[0].text = "Magnitud"; rt.rows[0].cells[1].text = "Valor"
    claves_preferidas = ["Q_m3s","iteraciones","residual_m","convergencia","total_hf_m","total_hm_accesorios_m","total_hm_transiciones_m","hL_total_m","D_mott","P2_calculada","margen_presion","satisfactorio"]
    vistos = set()
    for k in claves_preferidas:
        if k in result and not isinstance(result[k], (dict,list)):
            c = rt.add_row().cells; c[0].text = k; c[1].text = _fmt(result[k]); vistos.add(k)
    for k,v in result.items():
        if k not in vistos and not isinstance(v,(dict,list)):
            c = rt.add_row().cells; c[0].text = str(k); c[1].text = _fmt(v)

    filas = result.get("resultados_tramos") or []
    if filas:
        sec_land = doc.add_section(WD_SECTION.NEW_PAGE)
        sec_land.orientation = WD_ORIENT.LANDSCAPE
        sec_land.page_width, sec_land.page_height = portrait_height, portrait_width
        sec_land.left_margin = Inches(0.45); sec_land.right_margin = Inches(0.45)
        sec_land.top_margin = Inches(0.55); sec_land.bottom_margin = Inches(0.55)
        doc.add_heading("7.1 Resultados por tramo", level=2)
        columnas = [
            ("Tramo","Tramo"),("L (m)","L (m)"),("D (m)","D (m)"),("V (m/s)","V (m/s)"),
            ("Re","Re"),("Régimen","regimen"),("f Darcy","f Darcy"),("ΣK","K"),("hf (m)","hf (m)"),("hm (m)","hm (m)"),
        ]
        tabr = doc.add_table(rows=1, cols=len(columnas)); tabr.style = "Table Grid"; tabr.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i,(h,_) in enumerate(columnas): tabr.rows[0].cells[i].text = h
        for fila in filas:
            c = tabr.add_row().cells
            for i,(_,key) in enumerate(columnas): c[i].text = _fmt(fila.get(key))

    # Figuras: si hubo tabla de tramos, permanecen en la misma sección horizontal.
    doc.add_heading("8. Diagramas hidráulicos", level=1)
    figs = _figuras_png(contexto_diagramas)
    if figs:
        for idx,(titulo, png) in enumerate(figs,1):
            doc.add_heading(f"Figura {idx}. {titulo}", level=2)
            bio = io.BytesIO(png)
            doc.add_picture(bio, width=Inches(9.3))
            p = doc.paragraphs[-1]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc.add_paragraph("No se dispuso del contexto V14.15 necesario para incrustar figuras en esta ejecución.")

    # Regreso a formato vertical para auditorías y cierre.
    if filas:
        sec_portrait = doc.add_section(WD_SECTION.NEW_PAGE)
        sec_portrait.orientation = WD_ORIENT.PORTRAIT
        sec_portrait.page_width, sec_portrait.page_height = portrait_width, portrait_height
        sec_portrait.left_margin = Inches(0.7); sec_portrait.right_margin = Inches(0.7)
        sec_portrait.top_margin = Inches(0.65); sec_portrait.bottom_margin = Inches(0.65)

    # Auditorías
    doc.add_heading("9. Auditoría matemática V14.8.3", level=1)
    aud = r.get("auditoria_v148") or {}
    doc.add_paragraph(f"Estado: {aud.get('estado','—')} · errores: {aud.get('errores',0)} · advertencias: {aud.get('advertencias',0)}")
    for c0 in aud.get("checks") or []:
        doc.add_paragraph(f"[{str(c0.get('nivel','')).upper()}] {c0.get('codigo','')}: {c0.get('mensaje','')}", style="List Bullet")

    doc.add_heading("10. Auditoría física V14.16", level=1)
    audf = r.get("auditoria_fisica_v1416") or {}
    if audf:
        doc.add_paragraph(f"Estado: {audf.get('estado','—')} · errores: {audf.get('errores',0)} · advertencias: {audf.get('advertencias',0)} · informativos: {audf.get('informativos',0)}")
        for c0 in audf.get("checks") or []:
            doc.add_paragraph(f"[{str(c0.get('nivel','')).upper()}] {c0.get('codigo','')}: {c0.get('mensaje','')}", style="List Bullet")
    else:
        doc.add_paragraph("No evaluada.")

    doc.add_page_break()
    doc.add_heading("11. Nomenclatura", level=1)
    tn = doc.add_table(rows=1, cols=3); tn.style = "Table Grid"; tn.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(("Símbolo","Descripción","Unidad")): tn.rows[0].cells[i].text = h
    for a,b,c0 in _nomenclatura():
        cc = tn.add_row().cells; cc[0].text=a; cc[1].text=b; cc[2].text=c0

    doc.add_heading("12. Trazabilidad", level=1)
    doc.add_paragraph(f"Reporte técnico V2: {VERSION_REPORTE_V2}")
    doc.add_paragraph(f"Reporte base trazable: {r.get('version','V14.9')}")
    doc.add_paragraph(f"Firma V14.7: {r.get('firma_transferencia_v147','—')}")
    conf = r.get("confianza_datos_v1411") or {}
    if conf: doc.add_paragraph(f"Firma V14.11: {conf.get('firma','—')}")
    doc.add_paragraph("El reporte V14.18 presenta la información; no modifica el resultado del solver ni las auditorías.")

    # Encabezado/pie
    for s in doc.sections:
        hp = s.header.paragraphs[0]
        hp.text = f"Calculadora de sistemas de tuberías · {VERSION_REPORTE_V2}"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp = s.footer.paragraphs[0]
        fp.text = "Reporte generado automáticamente · verificar condiciones y supuestos antes de uso profesional"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.runs[0].font.size = Pt(7.5)

    out = io.BytesIO(); doc.save(out); return out.getvalue()


def reporte_pdf_v2(reporte: dict, contexto_diagramas: dict | None = None) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image, KeepTogether
    except ImportError as e:
        raise ErrorReporteV1418("Falta reportlab para generar el PDF V14.18.") from e

    r = reporte or {}
    bio = io.BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=A4, rightMargin=1.4*cm, leftMargin=1.4*cm, topMargin=1.5*cm, bottomMargin=1.5*cm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CenterTitle2", parent=styles["Title"], alignment=TA_CENTER, fontSize=18, leading=22, spaceAfter=12))
    story = []
    titulo_res, valor_res = _resultado_principal(r)
    story += [Spacer(1, 2.3*cm), Paragraph("REPORTE TÉCNICO DE SISTEMA DE TUBERÍAS", styles["CenterTitle2"]),
              Paragraph(str(r.get("clase") or "Sistema de tuberías"), styles["Heading2"]), Spacer(1, 0.6*cm)]
    data = [["Resultado principal", f"{titulo_res}: {valor_res}"], ["Auditoría matemática", str(r.get("estado_matematico") or "—")], ["Auditoría física", str(r.get("estado_fisico") or "—")], ["Snapshot V14.7", str(r.get("firma_transferencia_v147") or "—")[:12]]]
    tt = Table(data, colWidths=[5.0*cm, 10.0*cm]); tt.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.4,colors.grey),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("BACKGROUND",(0,0),(0,-1),colors.whitesmoke)])); story += [tt, PageBreak()]

    def h(txt, level=1): story.append(Paragraph(txt, styles["Heading1" if level==1 else "Heading2"]))
    def p(txt): story.append(Paragraph(str(txt).replace("&","&amp;").replace("<","&lt;").replace(">","&gt;"), styles["BodyText"])); story.append(Spacer(1,0.15*cm))

    h("1. Resumen técnico"); p(f"Se resolvió {r.get('clase','—')}. Resultado principal: {titulo_res} = {valor_res}. Auditoría matemática: {r.get('estado_matematico','—')}; auditoría física: {r.get('estado_fisico','—')}.")
    h("2. Enunciado"); p(r.get("enunciado") or "No disponible.")
    h("3. Datos y procedencia")
    rows = [["Dato","Valor","Unidad","Fuente"]]
    for x in (r.get("entradas") or {}).get("generales") or []:
        rows.append([str(x.get("dato","")), _fmt(x.get("valor")), str(x.get("unidad","")), str(x.get("fuente",""))])
    tb = Table(rows, repeatRows=1, colWidths=[4.0*cm,3.0*cm,2.2*cm,7.0*cm]); tb.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.3,colors.grey),("BACKGROUND",(0,0),(-1,0),colors.whitesmoke),("FONTSIZE",(0,0),(-1,-1),7.2)])); story += [tb, Spacer(1,0.3*cm)]

    h("4. Supuestos y convenciones")
    for s in _supuestos(r): p("• " + s)
    h("5. Ecuaciones y método")
    for nombre, eq in _ecuaciones_metodo(str(r.get("clase") or ""), r.get("resultado") or {}): p(f"{nombre}: {eq}")
    h("6. Resultados")
    rr = [["Magnitud","Valor"]]
    for k,v in (r.get("resultado") or {}).items():
        if not isinstance(v,(dict,list)): rr.append([str(k), _fmt(v)])
    tr = Table(rr, repeatRows=1, colWidths=[8*cm,8*cm]); tr.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.3,colors.grey),("BACKGROUND",(0,0),(-1,0),colors.whitesmoke)])); story += [tr, Spacer(1,0.3*cm)]
    filas_pdf = (r.get("resultado") or {}).get("resultados_tramos") or []
    if filas_pdf:
        h("6.1 Resultados por tramo", 2)
        rr2 = [["Tramo","D (m)","V (m/s)","Re","f Darcy","hf (m)","hm (m)"]]
        for fila in filas_pdf:
            rr2.append([_fmt(fila.get("Tramo")), _fmt(fila.get("D (m)")), _fmt(fila.get("V (m/s)")), _fmt(fila.get("Re")), _fmt(fila.get("f Darcy")), _fmt(fila.get("hf (m)")), _fmt(fila.get("hm (m)"))])
        tr2 = Table(rr2, repeatRows=1, colWidths=[1.5*cm,2.2*cm,2.2*cm,3.0*cm,2.2*cm,2.2*cm,2.2*cm])
        tr2.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.3,colors.grey),("BACKGROUND",(0,0),(-1,0),colors.whitesmoke),("FONTSIZE",(0,0),(-1,-1),7.2)]))
        story += [tr2, Spacer(1,0.3*cm)]

    figs = _figuras_png(contexto_diagramas)
    h("7. Diagramas hidráulicos")
    if figs:
        for idx,(tit,png) in enumerate(figs,1):
            p(f"Figura {idx}. {tit}")
            im = Image(io.BytesIO(png)); im.drawWidth = 17.5*cm; im.drawHeight = im.imageHeight * (im.drawWidth / im.imageWidth)
            story += [im, Spacer(1,0.35*cm)]
    else:
        p("No se dispuso del contexto V14.15 para incrustar figuras.")

    h("8. Auditoría matemática V14.8.3")
    aud = r.get("auditoria_v148") or {}; p(f"Estado: {aud.get('estado','—')} · errores: {aud.get('errores',0)} · advertencias: {aud.get('advertencias',0)}")
    for c in aud.get("checks") or []: p(f"[{str(c.get('nivel','')).upper()}] {c.get('codigo','')}: {c.get('mensaje','')}")
    h("9. Auditoría física V14.16")
    audf = r.get("auditoria_fisica_v1416") or {}; p(f"Estado: {audf.get('estado','NO EVALUADA')}")
    for c in audf.get("checks") or []: p(f"[{str(c.get('nivel','')).upper()}] {c.get('codigo','')}: {c.get('mensaje','')}")
    h("10. Trazabilidad"); p(f"V14.18 sobre reporte base {r.get('version','V14.9')} · firma V14.7: {r.get('firma_transferencia_v147','—')}")

    doc.build(story)
    return bio.getvalue()
