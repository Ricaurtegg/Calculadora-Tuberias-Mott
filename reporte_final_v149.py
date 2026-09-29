"""Trazabilidad de ejecución y reporte técnico final — V14.9.

V14.9 no altera el solver hidráulico. Une las capas ya validadas:
- V14.5 expediente consolidado;
- V14.6 respuestas faltantes;
- V14.7 snapshot confirmado antes de transferir;
- V14.8.3 auditoría independiente posterior al solver.

Produce un manifiesto inmutable de entrada y reportes Markdown, JSON y DOCX.
"""
from __future__ import annotations

import copy
import io
import json
from datetime import datetime, timezone

from consolidacion_datos import consolidar_resultado
from transferencia_v147 import construir_previsualizacion_transferencia

VERSION_REPORTE = "V14.9"
INTEGRACION_V1412_MANIFIESTO = "V14.12.1"
INTEGRACION_V1413_MANIFIESTO = "V14.13"
INTEGRACION_V1414_MANIFIESTO = "V14.14"


class ErrorReporteV149(ValueError):
    pass


def _ahora_utc():
    return datetime.now(timezone.utc).isoformat()


def _fuentes_respuestas(resultado: dict):
    mapa = {}
    for r in (resultado or {}).get("respuestas_v146") or []:
        clave = str(r.get("clave") or "")
        tramo = r.get("tramo")
        mapa[(clave, tramo)] = "usuario V14.6"
    return mapa


def _fuente_general(clave: str, resultado: dict, expediente: dict, respuestas: dict):
    inc = {str(x.get("clave")) for x in expediente.get("incognitas") or []}
    if clave in inc:
        return "incógnita del problema"
    if (clave, None) in respuestas:
        return respuestas[(clave, None)]
    inf = {str(x.get("clave")) for x in expediente.get("inferencias") or []}
    if clave in inf:
        return "inferencia segura V14.5"
    return "texto/OCR/figura confirmada"


def construir_manifiesto_ejecucion(resultado: dict, enunciado: str = "", previsualizacion: dict | None = None) -> dict:
    base = consolidar_resultado(copy.deepcopy(resultado or {}), enunciado)
    preview = copy.deepcopy(previsualizacion) if previsualizacion else construir_previsualizacion_transferencia(base, enunciado)
    if not preview.get("listo_para_transferir"):
        raise ErrorReporteV149("V14.9 solo crea el manifiesto cuando V14.7 está listo para transferir.")

    exp = base.get("expediente_v145") or {}
    prefill = base.get("prefill") or {}
    respuestas = _fuentes_respuestas(base)
    inc = {str(x.get("clave")) for x in exp.get("incognitas") or []}

    # V14.12.1 — conserva la trazabilidad del parser aunque una capa intermedia
    # haya promovido las normalizaciones desde prefill al nivel raíz. Esto hace
    # compatible el manifiesto con expedientes creados por V14.12 y con snapshots
    # históricos reconsolidados por V14.5/V14.11.
    normalizaciones_v1412 = copy.deepcopy(
        prefill.get("normalizaciones_v1412")
        or base.get("normalizaciones_v1412")
        or ((resultado or {}).get("prefill") or {}).get("normalizaciones_v1412")
        or (resultado or {}).get("normalizaciones_v1412")
        or []
    )
    parser_v1412 = copy.deepcopy(
        prefill.get("parser_v1412")
        or base.get("parser_v1412")
        or ((resultado or {}).get("prefill") or {}).get("parser_v1412")
        or (resultado or {}).get("parser_v1412")
        or {}
    )
    interpretacion_v1413 = copy.deepcopy(
        prefill.get("interpretacion_componentes_v1413")
        or base.get("interpretacion_componentes_v1413")
        or ((resultado or {}).get("prefill") or {}).get("interpretacion_componentes_v1413")
        or (resultado or {}).get("interpretacion_componentes_v1413")
        or {}
    )
    geometria_raw_v1414 = copy.deepcopy(
        base.get("geometria_v14_14")
        or (resultado or {}).get("geometria_v14_14")
        or {}
    )
    # Resumen trazable sin incrustar la imagen PNG anotada dentro del manifiesto.
    geometria_v1414 = {}
    if geometria_raw_v1414:
        geometria_v1414 = {
            "version": geometria_raw_v1414.get("version", "V14.14"),
            "confianza_global": geometria_raw_v1414.get("confianza_global"),
            "perspectiva": copy.deepcopy(geometria_raw_v1414.get("perspectiva") or {}),
            "calidad_imagen": copy.deepcopy(geometria_raw_v1414.get("calidad_imagen") or {}),
            "recorrido_polilinea": copy.deepcopy(geometria_raw_v1414.get("recorrido_polilinea") or {}),
            "niveles_relativos": copy.deepcopy(geometria_raw_v1414.get("niveles_relativos") or []),
            "cambios_nivel": copy.deepcopy(geometria_raw_v1414.get("cambios_nivel") or []),
            "orden": copy.deepcopy(geometria_raw_v1414.get("orden") or []),
            "depositos_candidatos": len(geometria_raw_v1414.get("depositos") or []),
            "equipos_candidatos": len(geometria_raw_v1414.get("equipos") or []),
            "cambios_seccion_candidatos": len(geometria_raw_v1414.get("cambios_seccion") or []),
            "advertencias": copy.deepcopy(geometria_raw_v1414.get("advertencias") or []),
        }

    generales = []
    for clave, etiqueta, unidad in (
        ("Q_m3s", "Caudal Q", "m³/s"),
        ("P1_kpa", "Presión P1", "kPa"),
        ("P2_kpa", "Presión P2", "kPa"),
        ("z1_m", "Elevación z1", "m"),
        ("z2_m", "Elevación z2", "m"),
        ("hA_m", "Carga de bomba hA", "m"),
        ("hR_m", "Carga de turbina hR", "m"),
    ):
        valor = "INCÓGNITA" if clave in inc else prefill.get(clave)
        generales.append({
            "clave": clave, "dato": etiqueta, "valor": valor, "unidad": unidad,
            "fuente": _fuente_general(clave, base, exp, respuestas),
        })

    fluido = prefill.get("fluido_app") or prefill.get("fluido_detectado")
    if not fluido:
        try:
            if (bool(prefill.get("propiedades_explicitas_v1412"))
                    and float(prefill.get("rho_usuario")) > 0
                    and float(prefill.get("nu_usuario_m2s")) > 0):
                fluido = "Propiedades explícitas (ρ/ν)"
        except (TypeError, ValueError):
            pass
    generales.insert(0, {
        "clave": "fluido_app", "dato": "Fluido", "valor": fluido, "unidad": "",
        "fuente": ("propiedades físicas explícitas" if fluido == "Propiedades explícitas (ρ/ν)"
                   else respuestas.get(("fluido_app", None), "texto/OCR/figura confirmada")),
    })

    tramos = []
    for i, t in enumerate(prefill.get("tramos") or [], 1):
        n = int(t.get("numero", i) or i)
        tramos.append({
            "numero": n,
            "L_m": t.get("L_m"),
            "D_m": "INCÓGNITA" if "D_m" in inc else t.get("D_m"),
            "material": t.get("material"),
            "epsilon_m": t.get("epsilon_m"),
            "accesorios": copy.deepcopy(t.get("accesorios") or []),
            "K_extra": t.get("K_extra"),
            "curvas": copy.deepcopy(t.get("curvas") or []),
            "fuentes": {
                "L_m": respuestas.get(("L_m", n), "texto/OCR/figura confirmada"),
                "D_m": "incógnita del problema" if "D_m" in inc else respuestas.get(("D_m", n), "texto/OCR/figura confirmada"),
                "material": respuestas.get(("material", n), "texto/OCR/figura confirmada"),
            },
        })

    return {
        "version": VERSION_REPORTE,
        "generado_utc": _ahora_utc(),
        "clase": base.get("clase"),
        "confianza_clase": base.get("confianza", 0),
        "enunciado": enunciado or "",
        "firma_transferencia_v147": preview.get("firma"),
        "firma_corta": preview.get("firma_corta"),
        "snapshot_v147": copy.deepcopy(preview.get("snapshot") or {}),
        "generales": generales,
        "tramos": tramos,
        "transiciones": copy.deepcopy(prefill.get("transiciones") or []),
        "inferencias_v145": copy.deepcopy(exp.get("inferencias") or []),
        "respuestas_v146": copy.deepcopy(base.get("respuestas_v146") or []),
        "confianza_datos_v1411": copy.deepcopy(base.get("confianza_datos_v1411") or {}),
        "normalizaciones_v1412": normalizaciones_v1412,
        "interpretacion_componentes_v1413": interpretacion_v1413,
        "geometria_v1414": geometria_v1414,
        "trazabilidad": {
            "expediente": exp.get("version"),
            "completado": (base.get("completado_v146") or {}).get("version"),
            "confianza_datos": (base.get("confianza_datos_v1411") or {}).get("version"),
            "parser_unidades": (parser_v1412.get("version") if normalizaciones_v1412 else None),
            "integracion_parser_reporte": INTEGRACION_V1412_MANIFIESTO,
            "componentes_mott": (interpretacion_v1413.get("version") if interpretacion_v1413 else None),
            "integracion_componentes_reporte": INTEGRACION_V1413_MANIFIESTO,
            "vision_geometrica": (geometria_v1414.get("version") if geometria_v1414 else None),
            "integracion_vision_reporte": INTEGRACION_V1414_MANIFIESTO,
            "transferencia": preview.get("version"),
            "reporte": VERSION_REPORTE,
        },
    }


def _normalizar_resultado(clase: str, resultado_solver: dict) -> dict:
    r = copy.deepcopy(resultado_solver or {})
    if clase in ("Clase II-A", "Clase II-B", "Clase II-C"):
        sistema = r.get("sistema") or {}
        return {
            "Q_m3s": r.get("Q_final"),
            "metodo_calculo_principal": r.get("metodo_calculo_principal"),
            "metodo_solicitado": r.get("metodo_solicitado"),
            "iteraciones": r.get("iteraciones"),
            "residual_m": r.get("residual"),
            "convergencia": r.get("convergencia", True),
            "total_hf_m": sistema.get("total_hf"),
            "total_hm_accesorios_m": sistema.get("total_hm_accesorios", sistema.get("total_hm")),
            "total_hm_transiciones_m": sistema.get("total_hm_transiciones", 0.0),
            "hL_total_m": sistema.get("hL_total"),
            "resultados_tramos": copy.deepcopy(sistema.get("resultados") or []),
            "transiciones_resultados": copy.deepcopy(sistema.get("transiciones_resultados") or []),
        }
    if clase == "Clase I":
        sistema = r.get("sistema") or {}
        return {
            "incognita": r.get("incognita"), "valor": r.get("valor"), "unidad": r.get("unidad"),
            "residual_m": r.get("residual"), "Q_m3s": r.get("Q_m3s"),
            "total_hf_m": sistema.get("total_hf"),
            "total_hm_accesorios_m": sistema.get("total_hm_accesorios", sistema.get("total_hm")),
            "total_hm_transiciones_m": sistema.get("total_hm_transiciones", 0.0),
            "hL_total_m": sistema.get("hL_total"),
            "resultados_tramos": copy.deepcopy(sistema.get("resultados") or []),
        }
    if clase == "Clase III-A":
        claves = ("D_mott", "D_numerico_verificacion", "diferencia_porcentaje_vs_numerico", "V", "Re", "regimen", "f", "hf", "residual", "tamano_comercial")
        return {k: copy.deepcopy(r.get(k)) for k in claves}
    if clase == "Clase III-B":
        claves = ("P2_calculada", "margen_presion", "satisfactorio", "V", "Re", "regimen", "f", "hf", "hm", "hL")
        return {k: copy.deepcopy(r.get(k)) for k in claves}
    return r


def construir_reporte_final(manifiesto: dict, resultado_solver: dict, auditoria: dict, auditoria_fisica: dict | None = None) -> dict:
    if not manifiesto or manifiesto.get("version") != VERSION_REPORTE:
        raise ErrorReporteV149("Se requiere un manifiesto V14.9 válido.")
    clase = manifiesto.get("clase")
    return {
        "version": VERSION_REPORTE,
        "generado_utc": _ahora_utc(),
        "clase": clase,
        "confianza_clase": manifiesto.get("confianza_clase"),
        "firma_transferencia_v147": manifiesto.get("firma_transferencia_v147"),
        "enunciado": manifiesto.get("enunciado", ""),
        "entradas": {
            "generales": copy.deepcopy(manifiesto.get("generales") or []),
            "tramos": copy.deepcopy(manifiesto.get("tramos") or []),
            "transiciones": copy.deepcopy(manifiesto.get("transiciones") or []),
        },
        "trazabilidad": copy.deepcopy(manifiesto.get("trazabilidad") or {}),
        "inferencias_v145": copy.deepcopy(manifiesto.get("inferencias_v145") or []),
        "respuestas_v146": copy.deepcopy(manifiesto.get("respuestas_v146") or []),
        "confianza_datos_v1411": copy.deepcopy(manifiesto.get("confianza_datos_v1411") or {}),
        "normalizaciones_v1412": copy.deepcopy(manifiesto.get("normalizaciones_v1412") or []),
        "interpretacion_componentes_v1413": copy.deepcopy(manifiesto.get("interpretacion_componentes_v1413") or {}),
        "geometria_v1414": copy.deepcopy(manifiesto.get("geometria_v1414") or {}),
        "resultado": _normalizar_resultado(str(clase), resultado_solver),
        "auditoria_v148": copy.deepcopy(auditoria or {}),
        "auditoria_fisica_v1416": copy.deepcopy(auditoria_fisica or {}),
        "estado_matematico": (auditoria or {}).get("estado", "SIN AUDITORÍA"),
        "estado_fisico": (auditoria_fisica or {}).get("estado", "NO EVALUADA") if auditoria_fisica else "NO EVALUADA",
        "estado_final": (
            "ERROR" if ((auditoria or {}).get("estado") == "ERROR" or (auditoria_fisica or {}).get("estado") == "ERROR")
            else "REVISAR" if ((auditoria or {}).get("estado") == "REVISAR" or (auditoria_fisica or {}).get("estado") == "REVISAR")
            else (auditoria or {}).get("estado", "SIN AUDITORÍA")
        ),
    }


def _fmt(v):
    if v is None:
        return "—"
    if isinstance(v, float):
        return f"{v:.8g}"
    return str(v)


def reporte_markdown(reporte: dict) -> str:
    r = reporte or {}
    lineas = [
        "# Reporte técnico — Sistema de tuberías en serie",
        "",
        f"**Versión:** {r.get('version','—')}",
        f"**Clase:** {r.get('clase','—')}",
        f"**Confianza de clasificación:** {r.get('confianza_clase','—')} %",
        f"**Estado de auditoría:** {r.get('estado_final','—')}",
        f"**Firma V14.7:** {r.get('firma_transferencia_v147','—')}",
        "",
        "## Enunciado",
        r.get("enunciado") or "—",
        "",
        "## Datos transferidos",
        "| Dato | Valor | Unidad | Fuente |",
        "|---|---:|---|---|",
    ]
    for x in (r.get("entradas") or {}).get("generales") or []:
        lineas.append(f"| {x.get('dato','')} | {_fmt(x.get('valor'))} | {x.get('unidad','')} | {x.get('fuente','')} |")

    lineas += ["", "## Tramos"]
    for t in (r.get("entradas") or {}).get("tramos") or []:
        lineas.append(f"### Tramo {t.get('numero','—')}")
        lineas.append(f"- L = {_fmt(t.get('L_m'))} m")
        lineas.append(f"- D = {_fmt(t.get('D_m'))} m")
        lineas.append(f"- Material = {_fmt(t.get('material'))}")
        acc = t.get("accesorios") or []
        lineas.append("- Accesorios: " + (", ".join(f"{a.get('cantidad',1)}× {a.get('nombre','Accesorio')}" for a in acc) if acc else "Ninguno"))

    conf = r.get("confianza_datos_v1411") or {}
    if conf:
        lineas += [
            "", "## Confianza individual V14.11",
            f"**Firma:** {conf.get('firma','—')}",
            f"**Confianza mínima:** {conf.get('confianza_min_pct','—')} %",
            f"**Confianza media:** {conf.get('confianza_media_pct','—')} %",
            f"**Pendientes de confirmación:** {conf.get('pendientes_confirmacion',0)}",
            "",
            "| Dato | Valor | Confianza | Nivel | Fuente | Estado |",
            "|---|---|---:|---|---|---|",
        ]
        for x in conf.get("entradas") or []:
            cp = "—" if x.get("nivel") == "Método" else f"{float(x.get('confianza_pct',0) or 0):.1f} %"
            lineas.append(
                f"| {x.get('dato','')} | {x.get('valor_mostrado','')} | {cp} | {x.get('nivel','')} | {x.get('fuente','')} | {x.get('estado','')} |"
            )

    norm = r.get("normalizaciones_v1412") or []
    if norm:
        lineas += [
            "", "## Normalización de unidades y notación V14.12",
            "| Original | Valor SI | Unidad SI | Interpretación |",
            "|---|---:|---|---|",
        ]
        for e in norm:
            lineas.append(
                f"| {e.get('original','')} | {_fmt(e.get('valor_si'))} | {e.get('unidad_si','')} | {e.get('detalle','')} |"
            )

    comp13 = r.get("interpretacion_componentes_v1413") or {}
    if comp13.get("eventos"):
        lineas += ["", "## Interpretación de accesorios y transiciones V14.13"]
        for e in comp13.get("eventos") or []:
            if e.get("tipo") == "transicion":
                lineas.append(f"- Transición {e.get('entre')}→{int(e.get('entre',0))+1}: {e.get('detalle','')} · estado={e.get('estado','')}")
            else:
                tramo = e.get("tramo") if e.get("tramo") is not None else "confirmar"
                lineas.append(f"- {e.get('original','Accesorio')} · tramo {tramo} · {e.get('detalle','')}")

    geo14 = r.get("geometria_v1414") or {}
    if geo14:
        rec = geo14.get("recorrido_polilinea") or {}
        per = geo14.get("perspectiva") or {}
        cal = geo14.get("calidad_imagen") or {}
        lineas += [
            "", "## Evidencia geométrica V14.14",
            f"- Confianza geométrica global: {_fmt(geo14.get('confianza_global'))} %",
            f"- Segmentos del recorrido: {len(rec.get('segmentos') or [])}",
            f"- Niveles relativos detectados: {len(geo14.get('niveles_relativos') or [])}",
            f"- Cambios de sección candidatos: {geo14.get('cambios_seccion_candidatos',0)}",
            f"- Rectificación de perspectiva: {'sí' if per.get('aplicada') else 'no'}",
            f"- Calidad visual heurística: {cal.get('nivel','—')}",
            "",
            "> La evidencia geométrica describe posiciones y forma del recorrido; no convierte píxeles a metros ni sustituye cotas físicas.",
        ]

    lineas += ["", "## Resultado"]
    for k, v in (r.get("resultado") or {}).items():
        if isinstance(v, (list, dict)):
            continue
        lineas.append(f"- **{k}:** {_fmt(v)}")

    filas = (r.get("resultado") or {}).get("resultados_tramos") or []
    if filas:
        lineas += ["", "### Resultados por tramo"]
        cols = list(filas[0].keys())
        lineas.append("| " + " | ".join(cols) + " |")
        lineas.append("|" + "|".join(["---"] * len(cols)) + "|")
        for fila in filas:
            lineas.append("| " + " | ".join(_fmt(fila.get(c)) for c in cols) + " |")

    aud = r.get("auditoria_v148") or {}
    lineas += ["", "## Auditoría independiente V14.8.3", f"**Estado:** {aud.get('estado','—')}"]
    for c in aud.get("checks") or []:
        lineas.append(f"- [{str(c.get('nivel','')).upper()}] {c.get('codigo','')}: {c.get('mensaje','')}")

    audf = r.get("auditoria_fisica_v1416") or {}
    if audf:
        lineas += [
            "", "## Auditoría física ampliada V14.16",
            f"**Estado:** {audf.get('estado','—')} · errores={audf.get('errores',0)} · advertencias={audf.get('advertencias',0)} · informativos={audf.get('informativos',0)}",
            "",
            "> Los umbrales V14.16 son banderas de plausibilidad/revisión y no sustituyen criterios normativos específicos de diseño.",
        ]
        for c in audf.get("checks") or []:
            lineas.append(f"- [{str(c.get('nivel','')).upper()}] {c.get('codigo','')}: {c.get('mensaje','')}")

    inf = r.get("inferencias_v145") or []
    if inf:
        lineas += ["", "## Inferencias seguras V14.5"]
        for x in inf:
            lineas.append(f"- {x.get('clave','')}: {_fmt(x.get('valor'))} ({x.get('fuente','')})")
    resp = r.get("respuestas_v146") or []
    if resp:
        lineas += ["", "## Datos completados por el usuario V14.6"]
        for x in resp:
            lineas.append(f"- {x.get('clave','')}" + (f" — tramo {x.get('tramo')}" if x.get('tramo') else "") + f": {_fmt(x.get('respuesta'))}")

    lineas += ["", "---", "Generado por la calculadora académica basada en Mott 7.ª edición."]
    return "\n".join(lineas) + "\n"


def reporte_json(reporte: dict) -> str:
    return json.dumps(reporte, ensure_ascii=False, indent=2, default=str)


def reporte_docx(reporte: dict) -> bytes:
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError as e:
        raise ErrorReporteV149("Falta python-docx para generar el reporte Word.") from e

    r = reporte or {}
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
    doc.add_heading("Reporte técnico — Sistema de tuberías en serie", 0)
    p = doc.add_paragraph()
    p.add_run("Clase: ").bold = True; p.add_run(str(r.get("clase", "—")))
    p.add_run("    Auditoría: ").bold = True; p.add_run(str(r.get("estado_final", "—")))
    p.add_run("    V14.7: ").bold = True; p.add_run(str(r.get("firma_transferencia_v147", "—"))[:12])

    doc.add_heading("Enunciado", level=1)
    doc.add_paragraph(r.get("enunciado") or "—")

    doc.add_heading("Datos transferidos", level=1)
    gens = (r.get("entradas") or {}).get("generales") or []
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, h in enumerate(("Dato", "Valor", "Unidad", "Fuente")):
        table.rows[0].cells[i].text = h
    for x in gens:
        cells = table.add_row().cells
        cells[0].text = str(x.get("dato", "")); cells[1].text = _fmt(x.get("valor")); cells[2].text = str(x.get("unidad", "")); cells[3].text = str(x.get("fuente", ""))

    doc.add_heading("Tramos", level=1)
    for t in (r.get("entradas") or {}).get("tramos") or []:
        doc.add_heading(f"Tramo {t.get('numero','—')}", level=2)
        doc.add_paragraph(f"L = {_fmt(t.get('L_m'))} m; D = {_fmt(t.get('D_m'))} m; material = {_fmt(t.get('material'))}.")
        acc = t.get("accesorios") or []
        doc.add_paragraph("Accesorios: " + (", ".join(f"{a.get('cantidad',1)}× {a.get('nombre','Accesorio')}" for a in acc) if acc else "Ninguno"))

    conf = r.get("confianza_datos_v1411") or {}
    if conf:
        doc.add_heading("Confianza individual V14.11", level=1)
        doc.add_paragraph(
            f"Firma: {conf.get('firma','—')} · mínima: {conf.get('confianza_min_pct','—')} % · "
            f"media: {conf.get('confianza_media_pct','—')} % · pendientes: {conf.get('pendientes_confirmacion',0)}"
        )
        tabc = doc.add_table(rows=1, cols=6); tabc.style = "Table Grid"
        for i, h in enumerate(("Dato", "Valor", "Confianza", "Nivel", "Fuente", "Estado")):
            tabc.rows[0].cells[i].text = h
        for x in conf.get("entradas") or []:
            cells = tabc.add_row().cells
            cp = "—" if x.get("nivel") == "Método" else f"{float(x.get('confianza_pct',0) or 0):.1f} %"
            vals = (x.get("dato",""), x.get("valor_mostrado",""), cp, x.get("nivel",""), x.get("fuente",""), x.get("estado",""))
            for i, v in enumerate(vals): cells[i].text = str(v)

    norm = r.get("normalizaciones_v1412") or []
    if norm:
        doc.add_heading("Normalización de unidades y notación V14.12", level=1)
        tabn = doc.add_table(rows=1, cols=4); tabn.style = "Table Grid"
        for i, h in enumerate(("Original", "Valor SI", "Unidad SI", "Interpretación")):
            tabn.rows[0].cells[i].text = h
        for e in norm:
            cells = tabn.add_row().cells
            vals = (e.get("original",""), _fmt(e.get("valor_si")), e.get("unidad_si",""), e.get("detalle",""))
            for i, v in enumerate(vals): cells[i].text = str(v)

    comp13 = r.get("interpretacion_componentes_v1413") or {}
    if comp13.get("eventos"):
        doc.add_heading("Interpretación de accesorios y transiciones V14.13", level=1)
        for e in comp13.get("eventos") or []:
            if e.get("tipo") == "transicion":
                txt = f"Transición {e.get('entre')}→{int(e.get('entre',0))+1}: {e.get('detalle','')} · estado={e.get('estado','')}"
            else:
                tramo = e.get("tramo") if e.get("tramo") is not None else "confirmar"
                txt = f"{e.get('original','Accesorio')} · tramo {tramo} · {e.get('detalle','')}"
            doc.add_paragraph(txt, style="List Bullet")

    geo14 = r.get("geometria_v1414") or {}
    if geo14:
        rec = geo14.get("recorrido_polilinea") or {}
        per = geo14.get("perspectiva") or {}
        cal = geo14.get("calidad_imagen") or {}
        doc.add_heading("Evidencia geométrica V14.14", level=1)
        doc.add_paragraph(
            f"Confianza geométrica global: {_fmt(geo14.get('confianza_global'))} % · "
            f"segmentos del recorrido: {len(rec.get('segmentos') or [])} · "
            f"niveles relativos: {len(geo14.get('niveles_relativos') or [])} · "
            f"cambios de sección candidatos: {geo14.get('cambios_seccion_candidatos',0)}."
        )
        doc.add_paragraph(
            f"Rectificación de perspectiva: {'sí' if per.get('aplicada') else 'no'} · "
            f"calidad visual heurística: {cal.get('nivel','—')}."
        )
        doc.add_paragraph(
            "La evidencia geométrica describe posiciones y forma del recorrido; no convierte píxeles a metros ni sustituye cotas físicas."
        )

    doc.add_heading("Resultado", level=1)
    for k, v in (r.get("resultado") or {}).items():
        if not isinstance(v, (list, dict)):
            doc.add_paragraph(f"{k}: {_fmt(v)}")

    filas = (r.get("resultado") or {}).get("resultados_tramos") or []
    if filas:
        doc.add_heading("Resultados por tramo", level=2)
        cols = list(filas[0].keys())
        tab = doc.add_table(rows=1, cols=len(cols)); tab.style = "Table Grid"
        for i, c in enumerate(cols): tab.rows[0].cells[i].text = str(c)
        for fila in filas:
            cells = tab.add_row().cells
            for i, c in enumerate(cols): cells[i].text = _fmt(fila.get(c))

    aud = r.get("auditoria_v148") or {}
    doc.add_heading("Auditoría independiente V14.8.3", level=1)
    doc.add_paragraph(f"Estado: {aud.get('estado','—')} · comprobaciones: {aud.get('comprobaciones',0)} · errores: {aud.get('errores',0)} · advertencias: {aud.get('advertencias',0)}")
    for c in aud.get("checks") or []:
        doc.add_paragraph(f"[{str(c.get('nivel','')).upper()}] {c.get('codigo','')}: {c.get('mensaje','')}", style="List Bullet")

    audf = r.get("auditoria_fisica_v1416") or {}
    if audf:
        doc.add_heading("Auditoría física ampliada V14.16", level=1)
        doc.add_paragraph(
            f"Estado: {audf.get('estado','—')} · comprobaciones: {audf.get('comprobaciones',0)} · "
            f"errores: {audf.get('errores',0)} · advertencias: {audf.get('advertencias',0)} · informativos: {audf.get('informativos',0)}"
        )
        doc.add_paragraph("Los umbrales V14.16 son banderas de plausibilidad/revisión y no sustituyen criterios normativos específicos de diseño.")
        for c in audf.get("checks") or []:
            doc.add_paragraph(f"[{str(c.get('nivel','')).upper()}] {c.get('codigo','')}: {c.get('mensaje','')}", style="List Bullet")

    doc.add_heading("Trazabilidad", level=1)
    doc.add_paragraph(f"Versión del reporte: {VERSION_REPORTE}")
    doc.add_paragraph(f"Firma completa V14.7: {r.get('firma_transferencia_v147','—')}")
    doc.add_paragraph("La procedencia de cada dato se conserva en la tabla de datos transferidos.")
    if r.get("confianza_datos_v1411"):
        doc.add_paragraph(f"Firma V14.11: {(r.get('confianza_datos_v1411') or {}).get('firma','—')}")
    if r.get("normalizaciones_v1412"):
        doc.add_paragraph(f"V14.12 conservó {len(r.get('normalizaciones_v1412') or [])} normalización(es) de notación/unidades con su fragmento original.")
    if (r.get("interpretacion_componentes_v1413") or {}).get("eventos"):
        doc.add_paragraph(f"V14.13 conservó {len((r.get('interpretacion_componentes_v1413') or {}).get('eventos') or [])} evento(s) de interpretación de accesorios/transiciones.")
    if r.get("geometria_v1414"):
        doc.add_paragraph("V14.14 conservó un resumen trazable de la evidencia geométrica confirmada, sin incrustar ni convertir magnitudes de píxeles.")
    if r.get("auditoria_fisica_v1416"):
        doc.add_paragraph(f"V14.16 añadió auditoría física ampliada con estado {(r.get('auditoria_fisica_v1416') or {}).get('estado','—')}.")

    bio = io.BytesIO(); doc.save(bio); return bio.getvalue()


def verificar_ejecucion_vs_manifiesto(manifiesto: dict, entrada_actual: dict, tol_rel: float = 1e-8) -> dict:
    """Comprueba que los datos realmente ejecutados sigan coincidiendo con V14.7.

    Las incógnitas se omiten. Esto permite editar manualmente después de transferir,
    pero V14.9 lo hará visible y no presentará la ejecución como trazable.
    """
    m = manifiesto or {}
    a = entrada_actual or {}
    diferencias = []

    def iguales(x, y):
        if x is None and y is None:
            return True
        if isinstance(x, (int, float)) and isinstance(y, (int, float)):
            ref = max(abs(float(x)), abs(float(y)), 1.0)
            return abs(float(x) - float(y)) <= tol_rel * ref
        return str(x) == str(y)

    mapa_actual = a.get("generales") or a
    for x in m.get("generales") or []:
        if x.get("valor") == "INCÓGNITA":
            continue
        clave = x.get("clave")
        if clave not in mapa_actual:
            continue
        if not iguales(x.get("valor"), mapa_actual.get(clave)):
            diferencias.append({"dato": clave, "esperado": x.get("valor"), "ejecutado": mapa_actual.get(clave)})

    at = a.get("tramos") or []
    for i, t in enumerate(m.get("tramos") or []):
        if i >= len(at):
            diferencias.append({"dato": f"tramo_{i+1}", "esperado": "presente", "ejecutado": "ausente"})
            continue
        cur = at[i]
        for clave in ("L_m", "D_m", "material"):
            esp = t.get(clave)
            if esp == "INCÓGNITA":
                continue
            if clave in cur and not iguales(esp, cur.get(clave)):
                diferencias.append({"dato": f"tramo_{i+1}.{clave}", "esperado": esp, "ejecutado": cur.get(clave)})

    return {
        "version": VERSION_REPORTE,
        "coincide_snapshot_v147": not diferencias,
        "diferencias": diferencias,
        "cantidad_diferencias": len(diferencias),
    }
