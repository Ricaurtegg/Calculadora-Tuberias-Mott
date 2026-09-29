"""Pruebas V14.2 autocontenidas.

Comprueban que V14.2 conserva TXT/DOCX/PDF seleccionable de V14.1 y añade
PNG/JPG/JPEG + PDF escaneado con OCR.
"""
from pathlib import Path
from io import BytesIO
import unicodedata

from lectura_archivos import (
    EXTENSIONES_SOPORTADAS,
    ErrorLecturaArchivo,
    extraer_texto_archivo,
    ocr_disponible,
)

BASE = Path(__file__).resolve().parent
APP = (BASE / "app.py").read_text(encoding="utf-8")


def ok(cond, msg):
    if not cond:
        raise AssertionError(msg)


def normalizar(s: str) -> str:
    s = (s or "").lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return " ".join(s.split())


def crear_docx_prueba() -> bytes:
    from docx import Document

    buf = BytesIO()
    doc = Document()
    doc.add_paragraph("Problema de prueba DOCX")
    tabla = doc.add_table(rows=2, cols=2)
    tabla.cell(0, 0).text = "Variable"
    tabla.cell(0, 1).text = "Valor"
    tabla.cell(1, 0).text = "Diámetro"
    tabla.cell(1, 1).text = "75 mm"
    doc.save(buf)
    return buf.getvalue()


def crear_pdf_texto_simple(texto: str) -> bytes:
    texto_pdf = texto.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({texto_pdf}) Tj ET".encode("latin-1", errors="replace")
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    salida = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objetos, 1):
        offsets.append(len(salida))
        salida.extend(f"{i} 0 obj\n".encode("ascii"))
        salida.extend(obj)
        salida.extend(b"\nendobj\n")
    xref = len(salida)
    salida.extend(f"xref\n0 {len(objetos)+1}\n".encode("ascii"))
    salida.extend(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        salida.extend(f"{off:010d} 00000 n \n".encode("ascii"))
    salida.extend(
        f"trailer\n<< /Size {len(objetos)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii")
    )
    return bytes(salida)


def _fuente_grande(tamano=46):
    from PIL import ImageFont

    candidatos = [
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibri.ttf",
        r"C:\Windows\Fonts\segoeui.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "DejaVuSans.ttf",
    ]
    for candidato in candidatos:
        try:
            return ImageFont.truetype(candidato, tamano)
        except Exception:
            pass
    return ImageFont.load_default()


def crear_imagen_prueba(formato: str) -> bytes:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (1800, 950), "white")
    draw = ImageDraw.Draw(img)
    fuente = _fuente_grande()
    lineas = [
        "Problema de prueba OCR V14.2",
        "Agua a 20 C fluye desde un deposito.",
        "Longitud L = 120 m.",
        "Diametro interior D = 102.3 mm.",
        "Presion P1 = 85 kPa.",
        "Elevaciones: z1 = 18 m; z2 = 5 m.",
        "Determine el caudal Q.",
    ]
    y = 70
    for linea in lineas:
        draw.text((80, y), linea, fill="black", font=fuente)
        y += 105
    bio = BytesIO()
    if formato.upper() == "JPEG":
        img.save(bio, format="JPEG", quality=95)
    else:
        img.save(bio, format=formato)
    return bio.getvalue()


def crear_pdf_escaneado() -> bytes:
    from PIL import Image

    img = Image.open(BytesIO(crear_imagen_prueba("PNG"))).convert("RGB")
    bio = BytesIO()
    img.save(bio, "PDF", resolution=200.0)
    return bio.getvalue()


print("[1/11] Extensiones V14.2...")
for ext in (".txt", ".docx", ".pdf", ".png", ".jpg", ".jpeg"):
    ok(ext in EXTENSIONES_SOPORTADAS, f"Falta {ext}")
print("  OK")

print("[2/11] TXT UTF-8 y Windows-1252 siguen funcionando...")
r = extraer_texto_archivo("problema.txt", "Determine el caudal en una tubería de 75 mm.".encode("utf-8"))
ok("Determine el caudal" in r.texto and not r.ocr_aplicado, "Falló TXT UTF-8")
r = extraer_texto_archivo("problema.txt", "Presión = 150 kPa; diámetro = 75 mm".encode("cp1252"))
ok("Presión" in r.texto and "diámetro" in r.texto, "Falló TXT cp1252")
print("  OK")

print("[3/11] DOCX con párrafos y tabla sigue funcionando...")
r = extraer_texto_archivo("problema.docx", crear_docx_prueba())
ok("Problema de prueba DOCX" in r.texto and "75 mm" in r.texto, "Falló DOCX")
ok(not r.ocr_aplicado, "DOCX de texto no debe marcar OCR")
print("  OK")

print("[4/11] PDF con texto seleccionable sigue usando pypdf...")
r = extraer_texto_archivo("problema.pdf", crear_pdf_texto_simple("Problema de prueba PDF seleccionable"))
ok("Problema de prueba PDF seleccionable" in r.texto, "Falló PDF seleccionable")
ok(r.paginas == 1 and not r.ocr_aplicado, "PDF seleccionable no debe activar OCR")
ok("pypdf" in r.metodo.lower(), "PDF seleccionable debe usar pypdf")
print("  OK")

print("[5/11] OCR disponible...")
disponible, detalle = ocr_disponible()
ok(disponible, f"OCR no disponible: {detalle}")
print(f"  OK: {detalle}")

print("[6/11] PNG con OCR...")
r = extraer_texto_archivo("problema.png", crear_imagen_prueba("PNG"))
t = normalizar(r.texto)
ok("agua a 20 c" in t and "120 m" in t, "PNG: OCR incompleto")
ok("102.3 mm" in t or "102,3 mm" in t, "PNG: no reconoció diámetro")
ok(r.ocr_aplicado, "PNG debe marcar OCR")
print("  OK")

print("[7/11] JPG con OCR...")
r = extraer_texto_archivo("problema.jpg", crear_imagen_prueba("JPEG"))
t = normalizar(r.texto)
ok("85 kpa" in t, "JPG: no reconoció presión")
ok("18 m" in t and "5 m" in t, "JPG: no reconoció cotas")
ok(r.ocr_aplicado, "JPG debe marcar OCR")
print("  OK")

print("[8/11] JPEG con OCR...")
r = extraer_texto_archivo("problema.jpeg", crear_imagen_prueba("JPEG"))
t = normalizar(r.texto)
ok("determine el caudal q" in t, "JPEG: no reconoció incógnita")
ok(r.ocr_aplicado, "JPEG debe marcar OCR")
print("  OK")

print("[9/11] PDF escaneado activa OCR automáticamente...")
r = extraer_texto_archivo("escaneado.pdf", crear_pdf_escaneado())
t = normalizar(r.texto)
ok("problema de prueba ocr v14.2" in t, "PDF escaneado: OCR no reconoció encabezado")
ok("determine el caudal q" in t, "PDF escaneado: OCR no reconoció enunciado")
ok(r.ocr_aplicado, "PDF escaneado debe marcar OCR")
ok("ocr" in normalizar(r.metodo), "Método debe indicar OCR")
print("  OK")

print("[10/11] Formato no admitido se rechaza...")
try:
    extraer_texto_archivo("archivo.bmp", b"abc")
except ErrorLecturaArchivo:
    pass
else:
    raise AssertionError("BMP no debe estar admitido todavía")
print("  OK")

print("[11/11] Integración V14.2 en app.py...")
for token in (
    "PNG, JPG y JPEG",
    "OCR automáticamente",
    "EXTENSIONES_SOPORTADAS",
    "st.file_uploader(",
    "ocr_aplicado",
    "Texto del problema — revise y edite antes de analizar",
):
    ok(token in APP, f"Falta integración V14.2 en app.py: {token}")
print("  OK")

print("\nTODAS LAS PRUEBAS V14.2 PASARON CORRECTAMENTE.")
