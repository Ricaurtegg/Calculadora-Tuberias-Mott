"""V14.19 - Procedimiento académico Mott.

Genera un desarrollo reproducible a partir del resultado YA resuelto por el motor.
No vuelve a resolver el problema ni sustituye los métodos usados por el solver.
"""
from __future__ import annotations

import io
import math
from typing import Any

VERSION_PROCEDIMIENTO = "V14.19"
G = 9.81


class ErrorProcedimientoV1419(ValueError):
    pass


def _f(v: Any, nd: int = 7) -> str:
    if v is None:
        return "—"
    if isinstance(v, bool):
        return "Sí" if v else "No"
    try:
        x = float(v)
    except (TypeError, ValueError):
        return str(v)
    if not math.isfinite(x):
        return str(x)
    ax = abs(x)
    if ax != 0 and (ax < 1e-4 or ax >= 1e6):
        return f"{x:.5e}"
    return f"{x:.{nd}g}"


def _snapshot_prefill(manifiesto: dict | None) -> dict:
    return (((manifiesto or {}).get("snapshot_v147") or {}).get("prefill") or {})


def _fila_resultado(resultado: dict, numero: int) -> dict:
    sis = resultado.get("sistema") or {}
    filas = sis.get("resultados") or []
    for fila in filas:
        try:
            if int(fila.get("Tramo", -1)) == int(numero):
                return fila
        except Exception:
            pass
    if 1 <= int(numero) <= len(filas):
        return filas[int(numero)-1]
    return {}


def _datos_frontera(contexto: dict | None) -> dict:
    c = contexto or {}
    gamma = float(c.get("gamma", 0.0) or 0.0)
    if gamma <= 0:
        return {}
    P1 = float(c.get("P1_pa", 0.0) or 0.0)
    P2 = float(c.get("P2_pa", 0.0) or 0.0)
    z1 = float(c.get("z1_m", 0.0) or 0.0)
    z2 = float(c.get("z2_m", 0.0) or 0.0)
    V1 = float(c.get("V1_ms", 0.0) or 0.0)
    V2 = float(c.get("V2_ms", 0.0) or 0.0)
    hA = float(c.get("hA_m", 0.0) or 0.0)
    hR = float(c.get("hR_m", 0.0) or 0.0)
    hv1 = V1 * V1 / (2.0 * G)
    hv2 = V2 * V2 / (2.0 * G)
    H1 = P1 / gamma + z1 + hv1
    H2 = P2 / gamma + z2 + hv2
    disponible = H1 + hA - hR - H2
    return {
        "gamma": gamma, "P1_pa": P1, "P2_pa": P2, "z1_m": z1, "z2_m": z2,
        "V1_ms": V1, "V2_ms": V2, "hv1_m": hv1, "hv2_m": hv2,
        "hA_m": hA, "hR_m": hR, "H1_m": H1, "H2_m": H2,
        "carga_disponible_m": disponible,
    }


def _paso(n: int, titulo: str, descripcion: str = "", ecuacion: str | None = None,
          sustitucion: str | None = None, resultado: str | None = None) -> dict:
    return {
        "numero": int(n), "titulo": titulo, "descripcion": descripcion,
        "ecuacion": ecuacion, "sustitucion": sustitucion, "resultado": resultado,
    }


def _tabla_tramos(resultado: dict) -> list[dict]:
    filas = ((resultado.get("sistema") or {}).get("resultados") or [])
    return [dict(x) for x in filas]


def _pasos_tramos(resultado: dict, nu: float | None, q: float | None, inicio: int) -> list[dict]:
    pasos = []
    n = inicio
    for fila in _tabla_tramos(resultado):
        i = fila.get("Tramo", len(pasos)+1)
        D = float(fila.get("D (m)", 0.0) or 0.0)
        L = float(fila.get("L (m)", 0.0) or 0.0)
        eps = float(fila.get("epsilon (m)", 0.0) or 0.0)
        A = float(fila.get("A (m2)", fila.get("A (m²)", 0.0)) or 0.0)
        V = float(fila.get("V (m/s)", 0.0) or 0.0)
        Re = float(fila.get("Re", 0.0) or 0.0)
        f = float(fila.get("f Darcy", 0.0) or 0.0)
        K = float(fila.get("K", fila.get("ΣK", 0.0)) or 0.0)
        hf = float(fila.get("hf (m)", 0.0) or 0.0)
        hm = float(fila.get("hm (m)", 0.0) or 0.0)
        eps_rel = eps / D if D > 0 else 0.0
        vh = V * V / (2.0 * G)
        qtxt = _f(q) if q is not None else _f(A * V)
        desc = f"Tramo {i}: se evalúan geometría, velocidad, Reynolds, fricción de Darcy y pérdidas con los valores realmente usados por el solver."
        sust = (
            f"A=π({_f(D)})²/4={_f(A)} m²; V={qtxt}/{_f(A)}={_f(V)} m/s; "
            f"Re=({_f(V)})({_f(D)})/{_f(nu) if nu else 'ν'}={_f(Re)}; ε/D={_f(eps)}/{_f(D)}={_f(eps_rel)}; "
            f"f={_f(f)}; V²/(2g)={_f(vh)} m; "
            f"hf=({_f(f)})({_f(L)}/{_f(D)})({_f(vh)})={_f(hf)} m; "
            f"hm=({_f(K)})({_f(vh)})={_f(hm)} m."
        )
        pasos.append(_paso(n, f"Cálculo hidráulico del tramo {i}", desc,
                           "A=πD²/4; V=Q/A; Re=VD/ν; hf=f(L/D)(V²/2g); hm=K(V²/2g)",
                           sust, f"Tramo {i}: hf={_f(hf)} m; hm={_f(hm)} m; régimen={fila.get('regimen','—')}."))
        n += 1
    return pasos


def _historial_reducido(hist: list[dict]) -> list[dict]:
    """Conserva todo el historial: V14.19 pretende ser académico y auditable."""
    return [dict(x) for x in (hist or [])]


def construir_procedimiento_mott(
    clase: str,
    resultado_solver: dict,
    manifiesto: dict | None = None,
    propiedades_fluido: dict | None = None,
    contexto_diagramas: dict | None = None,
    auditoria: dict | None = None,
) -> dict:
    clase = str(clase or "")
    r = resultado_solver or {}
    if not clase or not r:
        raise ErrorProcedimientoV1419("Se requieren clase y resultado del solver.")

    prop = propiedades_fluido or {}
    nu = prop.get("nu")
    rho = prop.get("rho")
    gamma = prop.get("gamma") or ((float(rho) * G) if rho is not None else None)
    prefill = _snapshot_prefill(manifiesto)
    frontera = _datos_frontera(contexto_diagramas)
    pasos = []
    n = 1

    datos = {
        "clase": clase,
        "fluido": (prefill.get("fluido_app") or prefill.get("fluido_detectado")
                    or ("Propiedades explícitas (ρ/ν)" if (prefill.get("propiedades_explicitas_v1412")
                        and prefill.get("rho_usuario") is not None and prefill.get("nu_usuario_m2s") is not None) else "—")),
        "rho_kgm3": rho, "nu_m2s": nu, "gamma_Nm3": gamma,
        "P1_kpa": prefill.get("P1_kpa"), "P2_kpa": prefill.get("P2_kpa"),
        "z1_m": prefill.get("z1_m"), "z2_m": prefill.get("z2_m"),
        "hA_m": prefill.get("hA_m", 0.0), "hR_m": prefill.get("hR_m", 0.0),
        "Q_m3s": prefill.get("Q_m3s"),
    }
    pasos.append(_paso(n, "Datos y clasificación", f"El expediente confirmado V14.7 clasifica el problema como {clase}. Se usan únicamente datos confirmados y propiedades efectivas del solver.",
                       None,
                       f"ρ={_f(rho)} kg/m³; ν={_f(nu)} m²/s; γ={_f(gamma)} N/m³; P1={_f(prefill.get('P1_kpa'))} kPa; P2={_f(prefill.get('P2_kpa'))} kPa; z1={_f(prefill.get('z1_m'))} m; z2={_f(prefill.get('z2_m'))} m.",
                       f"Método seleccionado: {r.get('metodo_calculo_principal') or clase}."))
    n += 1

    if frontera:
        pasos.append(_paso(n, "Ecuación general de energía y carga disponible",
                           "Se establece el balance entre los puntos 1 y 2 antes de introducir las pérdidas del sistema.",
                           "P1/γ + z1 + V1²/(2g) + hA - hR - hL = P2/γ + z2 + V2²/(2g)",
                           f"H1={_f(frontera['P1_pa'])}/{_f(frontera['gamma'])}+{_f(frontera['z1_m'])}+{_f(frontera['hv1_m'])}={_f(frontera['H1_m'])} m; H2={_f(frontera['P2_pa'])}/{_f(frontera['gamma'])}+{_f(frontera['z2_m'])}+{_f(frontera['hv2_m'])}={_f(frontera['H2_m'])} m; carga neta=H1+hA-hR-H2={_f(frontera['carga_disponible_m'])} m.",
                           f"Carga neta disponible = {_f(frontera['carga_disponible_m'])} m."))
        n += 1

    q_final = r.get("Q_final", r.get("Q"))

    if clase == "Clase II-A":
        md = r.get("mott_directo") or {}
        metodo_txt = str(r.get("metodo_calculo_principal") or "")
        if "prueba y error" in metodo_txt.lower():
            hist = list(r.get("historial_iteracion") or [])
            pasos.append(_paso(
                n,
                "Prueba y error con Darcy-Weisbach",
                "El enunciado exige explícitamente un enfoque de prueba y error. La Ec. (11-3) se usa únicamente como estimación inicial; el resultado principal se obtiene iterando Q hasta cerrar la ecuación de Darcy-Weisbach/energía.",
                "R(Q)=P1/γ+z1+V1²/(2g)+hA-hR-hf(Q)-P2/γ-z2-V2²/(2g) → 0",
                f"Estimación Mott de referencia={_f(r.get('Q_mott_referencia'))} m³/s; registros iterativos conservados={len(hist)}.",
                f"Q final={_f(q_final)} m³/s; residual={_f(r.get('residual'))} m; convergencia={r.get('convergencia')}."
            ))
            n += 1
        elif md.get("aplicable"):
            D = ((_tabla_tramos(r) or [{}])[0]).get("D (m)")
            L = ((_tabla_tramos(r) or [{}])[0]).get("L (m)")
            eps = ((_tabla_tramos(r) or [{}])[0]).get("epsilon (m)")
            hf_disp = md.get("h_f_disponible")
            pasos.append(_paso(n, "Aplicación directa del método II-A de Mott",
                               "Para una tubería uniforme sin pérdidas menores, el resultado principal se obtiene con la Ec. (11-3) de Mott basada en Swamee-Jain; el solver numérico queda solo como verificación independiente.",
                               "Q = -0.965 D² √(gDhf/L) ln[ε/(3.7D) + 1.78ν/(D√(gDhf/L))]",
                               f"D={_f(D)} m; L={_f(L)} m; ε={_f(eps)} m; ν={_f(nu)} m²/s; hf={_f(hf_disp)} m; √(gDhf/L)={_f(md.get('raiz_gDhL'))}; argumento ln={_f(md.get('argumento_log'))}.",
                               f"Q_Mott={_f(q_final)} m³/s; Q_verificación={_f(md.get('Q_numerico_verificacion'))} m³/s; diferencia={_f(md.get('diferencia_porcentaje_vs_numerico'))} %."))
            n += 1
    elif clase in ("Clase II-B", "Clase II-C"):
        q_iia = r.get("Q_IIA")
        pasos.append(_paso(n, "Estimación inicial y solución iterativa",
                           "Se parte de la estimación II-A cuando está disponible y se corrige hasta cerrar la ecuación de energía con las pérdidas correspondientes a la clase.",
                           "R(Q)=P1/γ+z1+V1²/(2g)+hA-hR-hL(Q)-P2/γ-z2-V2²/(2g) → 0",
                           f"Q inicial II-A={_f(q_iia)} m³/s; iteraciones internas reportadas={int(r.get('iteraciones',0) or 0)}; registros conservados en historial={len(r.get('historial_iteracion') or [])}.",
                           f"Q final={_f(q_final)} m³/s; residual={_f(r.get('residual'))} m; convergencia={r.get('convergencia')}."))
        n += 1

    if clase in ("Clase II-A", "Clase II-B", "Clase II-C"):
        pasos.extend(_pasos_tramos(r, float(nu) if nu is not None else None, float(q_final) if q_final is not None else None, n))
        n += len(_tabla_tramos(r))
        sis = r.get("sistema") or {}
        pasos.append(_paso(n, "Suma de pérdidas y comprobación",
                           "Se suman pérdidas distribuidas, accesorios y transiciones usando exactamente los totales del solver.",
                           "hL,total = Σhf + Σhm,accesorios + Σhm,transiciones",
                           f"Σhf={_f(sis.get('total_hf'))} m; Σhm,acc={_f(sis.get('total_hm_accesorios', sis.get('total_hm')))} m; Σhm,trans={_f(sis.get('total_hm_transiciones',0.0))} m.",
                           f"hL,total={_f(sis.get('hL_total'))} m; residual={_f(r.get('residual'))} m."))
        n += 1

    elif clase == "Clase I":
        sis = r.get("sistema") or {}
        q = r.get("Q_m3s")
        pasos.extend(_pasos_tramos(r, float(nu) if nu is not None else None, float(q) if q is not None else None, n))
        n += len(_tabla_tramos(r))
        inc = str(r.get("incognita") or "")
        val = r.get("valor")
        finales = r.get("estado_final") or {}
        pasos.append(_paso(n, f"Despeje de {r.get('incognita_etiqueta') or inc}",
                           "Con Q conocido se calculan primero las pérdidas; luego se despeja literalmente la incógnita indicada en la ecuación general de energía.",
                           "P1/γ + z1 + V1²/(2g) + hA - hR - hL = P2/γ + z2 + V2²/(2g)",
                           f"hL={_f(sis.get('hL_total'))} m; P1={_f(finales.get('P1'))} Pa; P2={_f(finales.get('P2'))} Pa; z1={_f(finales.get('z1'))} m; z2={_f(finales.get('z2'))} m; hA={_f(finales.get('hA'))} m; hR={_f(finales.get('hR'))} m.",
                           f"{r.get('incognita_etiqueta') or inc}={_f(val)} {r.get('unidad') or ''}; residual={_f(r.get('residual'))} m."))
        n += 1

    elif clase == "Clase III-A":
        D = r.get("D_mott", r.get("D_minimo"))
        pasos.append(_paso(n, "Diámetro mínimo mediante Mott Ec. (11-8)",
                           "Con Q y carga disponible conocidos, el diámetro mínimo se obtiene con la ecuación explícita III-A y se verifica independientemente con Colebrook.",
                           "D = 0.66 [ ε^1.25 (LQ²/(g hL))^4.75 + ν Q^9.4 (L/(g hL))^5.2 ]^0.04",
                           f"L/(ghL)={_f(r.get('L_sobre_ghL'))}; término rugosidad={_f(r.get('termino_rugosidad'))}; término viscosidad={_f(r.get('termino_viscosidad'))}; argumento={_f(r.get('argumento'))}.",
                           f"D_Mott={_f(D)} m; D_numérico={_f(r.get('D_numerico_verificacion'))} m; diferencia={_f(r.get('diferencia_porcentaje_vs_numerico'))} %."))
        n += 1
        pasos.append(_paso(n, "Verificación hidráulica del diámetro",
                           "El diámetro Mott se sustituye de nuevo en continuidad, Reynolds y Darcy-Weisbach.",
                           "A=πD²/4; V=Q/A; Re=VD/ν; hf=f(L/D)(V²/2g)",
                           f"A={_f(r.get('A'))} m²; V={_f(r.get('V'))} m/s; Re={_f(r.get('Re'))}; f={_f(r.get('f'))}; hf={_f(r.get('hf'))} m.",
                           f"Residual hf-hL,permitida={_f(r.get('residual'))} m; verificación numérica={_f(r.get('residual_numerico_verificacion'))} m."))
        n += 1

    elif clase == "Clase III-B":
        pasos.append(_paso(n, "Evaluación del diámetro disponible",
                           "Con D y Q conocidos se calculan las pérdidas del tramo y luego la presión realmente disponible en el punto 2.",
                           "A=πD²/4; V=Q/A; Re=VD/ν; hL=hf+hm",
                           f"D={_f(r.get('D'))} m; A={_f(r.get('A'))} m²; V={_f(r.get('V'))} m/s; Re={_f(r.get('Re'))}; f={_f(r.get('f'))}; hf={_f(r.get('hf'))} m; hm={_f(r.get('hm'))} m.",
                           f"hL={_f(r.get('hL'))} m."))
        n += 1
        pasos.append(_paso(n, "Presión disponible y criterio de cumplimiento",
                           "Se despeja P2 de la ecuación general de energía y se compara con la presión mínima requerida.",
                           "P2 = γ[P1/γ + z1 + V1²/(2g) + hA - hR - hL - z2 - V2²/(2g)]",
                           f"P2,calculada={_f(r.get('P2_calculada'))} Pa; P2,requerida={_f(r.get('P2_deseada'))} Pa.",
                           f"Margen={_f(r.get('margen_presion'))} Pa; cumple={'Sí' if bool(r.get('satisfactorio', r.get('cumple'))) else 'No'}."))
        n += 1

    estado_aud = str((auditoria or {}).get("estado") or "—")
    resultado_final = ""
    if clase in ("Clase II-A", "Clase II-B", "Clase II-C"):
        resultado_final = f"Q = {_f(q_final)} m³/s"
    elif clase == "Clase I":
        resultado_final = f"{r.get('incognita_etiqueta') or r.get('incognita')} = {_f(r.get('valor'))} {r.get('unidad') or ''}".strip()
    elif clase == "Clase III-A":
        resultado_final = f"D mínimo = {_f(r.get('D_mott', r.get('D_minimo')))} m"
    elif clase == "Clase III-B":
        resultado_final = f"P2 calculada = {_f(r.get('P2_calculada'))} Pa; margen = {_f(r.get('margen_presion'))} Pa"

    pasos.append(_paso(n, "Respuesta final y verificación", "El valor reportado es exactamente el producido por el solver; V14.19 no vuelve a resolver ni redondea internamente el problema.",
                       None, f"Auditoría matemática={estado_aud}.", resultado_final))

    return {
        "version": VERSION_PROCEDIMIENTO,
        "clase": clase,
        "metodo": r.get("metodo_calculo_principal") or clase,
        "firma_v147": (manifiesto or {}).get("firma_transferencia_v147"),
        "datos": datos,
        "frontera": frontera,
        "pasos": pasos,
        "resultados_tramos": _tabla_tramos(r),
        "historial_iteracion": _historial_reducido(r.get("historial_iteracion") or []),
        "resultado_final": resultado_final,
        "estado_auditoria": estado_aud,
        "trazabilidad": "Construido a partir del snapshot V14.7 y del resultado final del solver; no recalcula el problema.",
    }


def procedimiento_markdown(proc: dict) -> str:
    p = proc or {}
    out = [f"# Procedimiento Mott - {p.get('clase','—')}", "", f"**Versión:** {p.get('version','—')}", f"**Método:** {p.get('metodo','—')}", ""]
    for paso in p.get("pasos") or []:
        out.append(f"## {paso.get('numero')}. {paso.get('titulo')}")
        if paso.get("descripcion"): out.append(str(paso["descripcion"]))
        if paso.get("ecuacion"): out.extend(["", f"**Ecuación:** `{paso['ecuacion']}`"])
        if paso.get("sustitucion"): out.extend(["", f"**Sustitución:** {paso['sustitucion']}"])
        if paso.get("resultado"): out.extend(["", f"**Resultado:** {paso['resultado']}"])
        out.append("")
    hist = p.get("historial_iteracion") or []
    if hist:
        out.extend(["## Historial de iteración", "", "| Iter. | Q (m³/s) | hL (m) | Residual (m) | Re mín | Re máx | f mín | f máx |", "|---:|---:|---:|---:|---:|---:|---:|---:|"])
        for row in hist:
            out.append("| {i} | {q} | {hl} | {res} | {rmin} | {rmax} | {fmin} | {fmax} |".format(
                i=row.get("Iteración","—"), q=_f(row.get("Q (m³/s)")), hl=_f(row.get("hL total (m)")), res=_f(row.get("Residual (m)")),
                rmin=_f(row.get("Re mín")), rmax=_f(row.get("Re máx")), fmin=_f(row.get("f mín")), fmax=_f(row.get("f máx"))))
        out.append("")
    out.extend(["## Respuesta final", "", f"**{p.get('resultado_final','—')}**", "", f"Trazabilidad: {p.get('trazabilidad','')}"])
    return "\n".join(out)


def procedimiento_docx(proc: dict) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.section import WD_SECTION, WD_ORIENT
    except ImportError as e:
        raise ErrorProcedimientoV1419("Falta python-docx para generar el procedimiento Word.") from e

    p = proc or {}
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65); sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)
    doc.styles["Normal"].font.name = "Aptos"; doc.styles["Normal"].font.size = Pt(10)
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"PROCEDIMIENTO MOTT - {p.get('clase','—')}"); run.bold = True; run.font.size = Pt(18)
    sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"{p.get('metodo','—')} · {VERSION_PROCEDIMIENTO}").italic = True

    for paso in p.get("pasos") or []:
        doc.add_heading(f"{paso.get('numero')}. {paso.get('titulo')}", level=1)
        if paso.get("descripcion"): doc.add_paragraph(str(paso["descripcion"]))
        for rotulo, clave in (("Ecuación", "ecuacion"), ("Sustitución", "sustitucion"), ("Resultado", "resultado")):
            if paso.get(clave):
                par = doc.add_paragraph(); par.add_run(rotulo + ": ").bold = True; par.add_run(str(paso[clave]))

    hist = p.get("historial_iteracion") or []
    if hist:
        portrait_width, portrait_height = sec.page_width, sec.page_height
        sec_land = doc.add_section(WD_SECTION.NEW_PAGE)
        sec_land.orientation = WD_ORIENT.LANDSCAPE
        sec_land.page_width, sec_land.page_height = portrait_height, portrait_width
        sec_land.left_margin = Inches(0.45); sec_land.right_margin = Inches(0.45)
        sec_land.top_margin = Inches(0.5); sec_land.bottom_margin = Inches(0.5)
        doc.add_heading("Historial de iteración", level=1)
        doc.add_paragraph(
            "La tabla conserva las filas registradas por el solver. El contador interno total puede ser mayor "
            "si V14.10 ejecutó refinamientos numéricos adicionales después del historial principal."
        )
        cols = [("Iter.","Iteración"),("Q (m³/s)","Q (m³/s)"),("hL (m)","hL total (m)"),("Residual (m)","Residual (m)"),("Re mín","Re mín"),("Re máx","Re máx"),("f mín","f mín"),("f máx","f máx")]
        t = doc.add_table(rows=1, cols=len(cols)); t.style="Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i,(h,_) in enumerate(cols):
            t.rows[0].cells[i].text=h
            t.rows[0].cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for row in hist:
            c=t.add_row().cells
            for i,(_,key) in enumerate(cols):
                c[i].text=_f(row.get(key))
                c[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        sec_portrait = doc.add_section(WD_SECTION.NEW_PAGE)
        sec_portrait.orientation = WD_ORIENT.PORTRAIT
        sec_portrait.page_width, sec_portrait.page_height = portrait_width, portrait_height
        sec_portrait.left_margin = Inches(0.75); sec_portrait.right_margin = Inches(0.75)
        sec_portrait.top_margin = Inches(0.65); sec_portrait.bottom_margin = Inches(0.65)

    doc.add_heading("Respuesta final", level=1)
    rr = doc.add_paragraph(); run = rr.add_run(str(p.get("resultado_final") or "—")); run.bold=True; run.font.size=Pt(12)
    doc.add_paragraph(str(p.get("trazabilidad") or ""))
    out=io.BytesIO(); doc.save(out); return out.getvalue()
